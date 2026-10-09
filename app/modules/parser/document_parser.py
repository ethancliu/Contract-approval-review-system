"""文档解析器。

把本地已存合同文件解析为统一的「全文 + 分段定位结构」：

    docx  → python-docx（段落 + 表格）
    pdf   → pypdf（按页提取文本层；纯扫描 PDF 无文本层时转 OCR 阻塞）
    txt   → 按编码尝试读取（utf-8 / utf-8-sig / gbk）
    图片  → OCR 引擎识别（见 ocr_engine.py）

异常处理：所有解析失败均抛出带有明确原因的 BlockedError 子类，
由上层置任务为 blocked 并记录阻塞原因，绝不静默返回空结果。
"""

from __future__ import annotations

from pathlib import Path

from ...constants import TEXT_DOC_EXTS
from ...exceptions import EmptyContentError, ParseError
from .ocr_engine import OcrEngine, ocr_engine

#: 图片扩展名
_IMAGE_EXTS = (".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp")


# ======================================================================
# 分段构造
# ======================================================================
def build_segments(raw_blocks: list[tuple[int, str, str]], text_source: str) -> list[dict]:
    """把 [(页码, 文本, 来源)] 转成带字符偏移的分段定位结构。

    参数：
        raw_blocks: 每个元素为 (页码, 段落文本, 来源标识)
        text_source: 来源标识（docx/pdf/txt/ocr）
    返回：
        分段列表，每段含 page/block/line/char_start/char_end/text/source
    """
    segments: list[dict] = []
    offset = 0
    block_no = 0

    for page, text, source in raw_blocks:
        text = (text or "").replace("\r\n", "\n").replace("\r", "\n")
        if not text.strip():
            continue
        block_no += 1
        lines = text.split("\n") if "\n" in text else [text]
        for line_no, line in enumerate(lines, start=1):
            if not line.strip():
                offset += len(line) + 1
                continue
            start = offset
            end = offset + len(line)
            segments.append({
                "page": page,
                "block": block_no,
                "line": line_no,
                "char_start": start,
                "char_end": end,
                "text": line.strip(),
                "source": source or text_source,
            })
            offset = end + 1  # +1 为拼接换行符

    return segments


def segments_to_text(segments: list[dict]) -> str:
    """把分段还原为全文。"""
    return "\n".join(seg.get("text", "") for seg in segments)


# ======================================================================
# 各格式解析实现
# ======================================================================
def _parse_txt(path: Path) -> tuple[list[tuple[int, str, str]], int]:
    content = None
    for encoding in ("utf-8", "utf-8-sig", "gbk", "gb18030"):
        try:
            content = path.read_text(encoding=encoding)
            break
        except UnicodeDecodeError:
            continue
    if content is None:
        raise ParseError(
            "文件解析异常：文本文件编码无法识别",
            detail=f"文件：{path.name}；已尝试编码 utf-8 / utf-8-sig / gbk / gb18030",
        )
    blocks = [(1, p, "txt") for p in content.split("\n\n") if p.strip()]
    if not blocks:
        blocks = [(1, content, "txt")]
    return blocks, 1


def _parse_docx(path: Path) -> tuple[list[tuple[int, str, str]], int]:
    try:
        import docx  # noqa: PLC0415
    except ImportError as exc:
        raise ParseError(
            "文件解析异常：缺少 docx 解析组件",
            detail=f"请执行 pip install python-docx（{exc}）",
        ) from exc

    try:
        document = docx.Document(str(path))
    except Exception as exc:  # noqa: BLE001 - 损坏文件 / 加密文件 / 非法 zip
        raise ParseError(
            "文件解析异常：Word 文档无法打开（文件可能已损坏或被加密）",
            detail=f"文件：{path.name}；错误：{exc}",
        ) from exc

    blocks: list[tuple[int, str, str]] = []
    for para in document.paragraphs:
        text = (para.text or "").strip()
        if text:
            blocks.append((1, text, "docx"))

    # 表格内容一并纳入（很多合同把金额、期限放在表格中）
    for table in document.tables:
        for row in table.rows:
            cells = [(c.text or "").strip().replace("\n", " ") for c in row.cells]
            line = " | ".join([c for c in cells if c])
            if line:
                blocks.append((1, line, "docx"))

    return blocks, 1


def _parse_pdf(path: Path) -> tuple[list[tuple[int, str, str]], int]:
    try:
        from pypdf import PdfReader  # noqa: PLC0415
    except ImportError as exc:
        raise ParseError(
            "文件解析异常：缺少 PDF 解析组件",
            detail=f"请执行 pip install pypdf（{exc}）",
        ) from exc

    try:
        reader = PdfReader(str(path))
    except Exception as exc:  # noqa: BLE001
        raise ParseError(
            "文件解析异常：PDF 文档无法打开（文件可能已损坏或被加密）",
            detail=f"文件：{path.name}；错误：{exc}",
        ) from exc

    blocks: list[tuple[int, str, str]] = []
    for page_no, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:  # noqa: BLE001 - 单页失败不影响其他页
            text = ""
        for para in text.split("\n\n"):
            if para.strip():
                blocks.append((page_no, para, "pdf"))

    if not blocks:
        # 无文本层 ⇒ 纯扫描 PDF，需 OCR；当前未接入 PDF 渲染组件，如实阻塞并给出处置建议
        raise EmptyContentError(
            "扫描件图片无法识别：PDF 无文本层（纯扫描件），未提取到任何文字",
            detail=f"文件：{path.name}；请将该 PDF 转为图片后重新上传，或由管理员接入 PDF 渲染 + OCR 流程后重试。",
        )

    return blocks, len(reader.pages)


def _parse_image(path: Path, engine: OcrEngine) -> tuple[list[tuple[int, str, str]], int]:
    result = engine.recognize(path, page=1)
    blocks = [(1, line, "ocr") for line in result["lines"]]
    return blocks, 1


# ======================================================================
# 统一入口
# ======================================================================
def parse_document(
    file_path: str | Path,
    engine: OcrEngine | None = None,
) -> dict:
    """解析本地合同文件。

    返回：
        {
          "file_path": str, "file_name": str, "file_type": str,
          "page_count": int, "segments": [...], "full_text": str,
          "parse_status": "success", "parse_error": "",
          "ocr_provider": ""            # 图片文件才有值
        }

    异常：解析/识别失败抛出 BlockedError 子类（由上层转为 blocked 状态）。
    """
    path = Path(file_path)
    if not path.exists():
        raise ParseError("文件解析异常：本地文件不存在", detail=str(path))

    ext = path.suffix.lower()
    engine = engine or ocr_engine

    if ext in _IMAGE_EXTS:
        blocks, page_count = _parse_image(path, engine)
        ocr_provider = engine._pick_engine(path)[0]  # noqa: SLF001 - 同一模块族内部使用
    elif ext == ".docx":
        blocks, page_count = _parse_docx(path)
        ocr_provider = ""
    elif ext == ".pdf":
        blocks, page_count = _parse_pdf(path)
        ocr_provider = ""
    elif ext in (".txt", ".md", ".text", ".json"):
        blocks, page_count = _parse_txt(path)
        ocr_provider = ""
    elif ext == ".doc":
        raise ParseError(
            "文件解析异常：不支持旧版 .doc 二进制格式",
            detail=f"文件：{path.name}；请另存为 .docx 或 .pdf 后重新提交审批。",
        )
    else:
        raise ParseError(
            "文件解析异常：不支持的文件格式",
            detail=f"文件：{path.name}；扩展名：{ext}；当前支持：{'、'.join(TEXT_DOC_EXTS)}、图片（扫描件）",
        )

    segments = build_segments(blocks, ext.lstrip("."))
    full_text = segments_to_text(segments)

    if not full_text.strip():
        raise EmptyContentError(
            "文档内容为空：解析后未提取到任何有效文本",
            detail=f"文件：{path.name}；页数/块数：{page_count}；段落数：{len(segments)}",
        )

    return {
        "file_path": str(path),
        "file_name": path.name,
        "file_type": ext.lstrip("."),
        "page_count": page_count,
        "segments": segments,
        "full_text": full_text,
        "char_count": len(full_text),
        "parse_status": "success",
        "parse_error": "",
        "ocr_provider": ocr_provider,
    }
